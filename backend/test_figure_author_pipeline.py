import os
import sys
import hashlib
import zipfile
import tempfile
import base64
import docx
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw

# Add backend directory to sys.path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app.models.udm import UniversalDocumentModel, Author, Affiliation
from app.models.template_spec import TemplateSpecification
from app.parsers.docx_parser import DocxParser
from app.renderers.latex_renderer import LatexRenderer
from app.compilation.latex_sandbox import LatexSandbox

def test_pipeline():
    temp_dir = tempfile.mkdtemp(prefix="test_pipeline_")
    print(f"Running automated test pipeline in {temp_dir}...")
    
    # -------------------------------------------------------------
    # 1. CREATE TEST DOCX WITH 3 DISTINCT FIGURES & MULTI-AUTHORS
    # -------------------------------------------------------------
    docx_path = os.path.join(temp_dir, "test_manuscript.docx")
    doc = docx.Document()
    
    # Title
    doc.add_heading("Multi-Figure and Multi-Author Test Manuscript", level=0)
    
    # Authors and Affiliations Line
    doc.add_paragraph("Alice Smith1, Bob Jones1,2, Charlie Brown2")
    doc.add_paragraph("1 Department of Computer Science, Stanford University, CA, USA")
    doc.add_paragraph("2 Department of Electrical Engineering, MIT, MA, USA")
    doc.add_paragraph("Abstract: Testing multi-figure extraction and author affiliation mapping.")
    
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph("First section paragraph with text.")
    
    # Generate 3 DISTINCT images with different colors & text
    fig_paths = []
    fig_hashes = []
    colors = [(220, 38, 38), (16, 185, 129), (37, 99, 235)]
    labels = ["Figure 1 Architecture", "Figure 2 Flowchart", "Figure 3 Results Chart"]
    
    for idx in range(3):
        img_p = os.path.join(temp_dir, f"source_fig_{idx+1}.png")
        img = Image.new('RGB', (300, 150), color=colors[idx])
        d = ImageDraw.Draw(img)
        d.text((30, 60), labels[idx], fill=(255, 255, 255))
        img.save(img_p)
        
        with open(img_p, "rb") as fh:
            img_bytes = fh.read()
            fig_hashes.append(hashlib.md5(img_bytes).hexdigest())
            
        doc.add_paragraph(f"Paragraph before figure {idx+1}")
        doc.add_picture(img_p, width=Inches(3))
        doc.add_paragraph(f"Figure {idx+1}: Caption for {labels[idx]}")
        
    doc.save(docx_path)
    print("Test DOCX created with 3 distinct images and 3 authors/2 affiliations.")
    
    # -------------------------------------------------------------
    # TEST A: DOCX PARSING INTO UDM
    # -------------------------------------------------------------
    udm = DocxParser.parse(docx_path)
    
    print("\n=== TEST A & B: UDM EXTRACTION VERIFICATION ===")
    print("Parsed Title:", udm.metadata.title)
    print("Authors Count:", len(udm.metadata.authors))
    for a in udm.metadata.authors:
        print(f"  Author: {a.name} -> Affil IDs: {a.affiliation_ids}")
        
    print("Affiliations Count:", len(udm.metadata.affiliations))
    for aff in udm.metadata.affiliations:
        print(f"  Affiliation [{aff.id}]: {aff.institution}")
        
    # Extract figures from UDM
    udm_figures = []
    for sec in udm.sections:
        for blk in sec.blocks:
            if blk.get("type") == "figure":
                udm_figures.append(blk)
                
    print(f"\nUDM Figures Extracted Count: {len(udm_figures)}")
    assert len(udm_figures) == 3, f"Expected 3 figures in UDM, got {len(udm_figures)}"
    
    udm_hashes = []
    for idx, fig in enumerate(udm_figures):
        b64 = fig.get("image_data_b64")
        assert b64 is not None, f"Figure {idx+1} missing image_data_b64"
        raw_b = base64.b64decode(b64)
        h = hashlib.md5(raw_b).hexdigest()
        udm_hashes.append(h)
        print(f"  Figure [{fig.get('id')}]: filename={fig.get('image_filename')}, caption='{fig.get('caption')}', hash={h[:8]}")
        
    assert len(set(udm_hashes)) == 3, "UDM figures do not have distinct image content hashes!"
    assert len(udm.metadata.authors) >= 3, "Failed to parse 3 authors"
    assert len(udm.metadata.affiliations) >= 2, "Failed to parse 2 affiliations"
    assert udm.metadata.authors[0].affiliation_ids == ["1"], "Author 1 affiliation mapping mismatch"
    assert udm.metadata.authors[1].affiliation_ids in [["1", "2"], ["1"], ["2"]], "Author 2 affiliation mapping mismatch"
    print("TEST A & TEST B PASSED SUCCESSFULLY!")
    
    # -------------------------------------------------------------
    # TEST C: RENDER INTO MULTIPLE LATEX TEMPLATE STYLES
    # -------------------------------------------------------------
    print("\n=== TEST C: MULTI-TEMPLATE AUTHOR & FIGURE RENDERING ===")
    
    template_styles = ["springer", "ieee", "elsevier", "lncs", "standard"]
    
    for style in template_styles:
        spec = TemplateSpecification(
            format_type="latex",
            document_class="sn-jnl" if style == "springer" else ("IEEEtran" if style == "ieee" else "article"),
            author_style=style
        )
        
        proj_out = os.path.join(temp_dir, f"proj_{style}")
        zip_out = os.path.join(temp_dir, f"proj_{style}.zip")
        
        created_files = LatexRenderer.render_project(
            udm=udm,
            spec=spec,
            dest_template_dir="",
            output_dir=proj_out,
            output_zip_path=zip_out
        )
        
        # Verify physical files written
        with open(os.path.join(proj_out, "main.tex"), "r", encoding="utf-8") as fh:
            tex_content = fh.read()
            
        print(f"\n--- Style: {style.upper()} ---")
        print("Generated TeX snippet (Authors & Figures):")
        for line in tex_content.split("\n"):
            if any(k in line for k in ["author", "affil", "address", "institute", "includegraphics"]):
                print("  ", line)
                
        # Assert distinct includegraphics
        inc_matches = [line for line in tex_content.split("\n") if "includegraphics" in line]
        assert len(inc_matches) == 3, f"Style {style}: Expected 3 includegraphics in main.tex, found {len(inc_matches)}"
        assert len(set(inc_matches)) == 3, f"Style {style}: includegraphics references are not distinct!"
        
        # TEST D: Verify Output ZIP contents
        with zipfile.ZipFile(zip_out, "r") as zf:
            zip_files = zf.namelist()
            fig_files = [f for f in zip_files if f.startswith("figures/")]
            assert len(fig_files) == 3, f"Style {style}: Expected 3 figure files in ZIP, found {len(fig_files)}"
            
            # Verify each zipped figure hash
            zip_hashes = [hashlib.md5(zf.read(f)).hexdigest() for f in fig_files]
            assert len(set(zip_hashes)) == 3, f"Style {style}: Zipped figure images are not distinct!"
            
        print(f"Style {style.upper()} verification PASSED!")
        
    # -------------------------------------------------------------
    # TEST E: LATEX SANDBOX COMPILATION TEST
    # -------------------------------------------------------------
    print("\n=== TEST E: LATEX SANDBOX COMPILATION ===")
    pdf_out = os.path.join(temp_dir, "preview.pdf")
    compiled, log = LatexSandbox.compile_project(
        project_dir=os.path.join(temp_dir, "proj_standard"),
        entrypoint="main.tex",
        udm=udm,
        output_pdf_path=pdf_out
    )
    print("PDF Compiled:", compiled)
    print("Log summary:", log[:200].replace("\n", " "))
    assert compiled is True, "PDF Compilation failed!"
    assert os.path.exists(pdf_out), "PDF file was not created!"
    print("TEST E COMPILATION PASSED!")
    
    print("\n==================================================")
    print("ALL REGRESSION TESTS (A, B, C, D, E) PASSED 100%!")
    print("==================================================")

if __name__ == "__main__":
    test_pipeline()
