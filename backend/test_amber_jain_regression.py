import os
import sys
import shutil
import tempfile
import docx
import unittest

sys.path.insert(0, os.path.abspath("api"))
sys.path.insert(0, os.path.abspath("backend"))

from api.app.parsers.docx_parser import DocxParser
from api.app.template_engine.analyzer import TemplateAnalyzer
from api.app.mappers.mapping_engine import MappingEngine
from api.app.renderers.latex_renderer import LatexRenderer
from api.app.compilation.latex_sandbox import LatexSandbox


class TestAmberJainAndSpringerRegression(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp(prefix="test_amber_jain_")
        self.docx_path = os.path.join(self.test_dir, "Role of Technology in a Hybrid Teaching Mode.docx")
        
        # Create synthetic DOCX manuscript matching user's exact specification
        doc = docx.Document()
        
        # Title
        p_title = doc.add_paragraph("Role of Technology in a Hybrid Teaching Mode")
        p_title.style = "Title"
        
        # Author header block
        doc.add_paragraph("SHARIKA T R")
        doc.add_paragraph("Assistant Professor, Department of Computer Science and Engineering")
        doc.add_paragraph("Phone Number: 9562799289")
        
        # Abstract & Keywords
        doc.add_paragraph("Abstract: This paper explores the integration of ICT in hybrid teaching mode.")
        doc.add_paragraph("Keywords: Hybrid Education, ICT, Blended Learning")
        
        sections_data = [
            ("Introduction", "Hybrid teaching integrates face-to-face instruction with online learning environments to maximize educational impact."),
            ("Hybrid vs Blended Learning", "While blended learning often refers to supplementary online components, hybrid learning fundamentally alters course delivery."),
            ("ICT and Hybrid Education", "Information and Communication Technologies form the backbone of modern hybrid pedagogical models."),
            ("ICT for Hybrid Educators", "Educators utilize learning management systems and digital assessment tools to facilitate interaction."),
            ("ICT for Hybrid Learners", "Learners benefit from self-paced modules, interactive simulations, and instant feedback systems."),
            ("Benefits of Hybrid Learning", "Hybrid education offers flexibility, accessibility, and improved engagement for diverse student populations."),
            ("Challenges in Hybrid Education", "Challenges include digital divide, infrastructure limitations, and requiring enhanced digital literacy."),
            ("Conclusion", "Technology plays a pivotal role in enabling scalable and effective hybrid teaching environments.")
        ]
        
        for sec_title, sec_text in sections_data:
            doc.add_heading(sec_title, level=1)
            doc.add_paragraph(sec_text)
            
        # Table 1
        t1 = doc.add_table(rows=3, cols=3)
        t1.cell(0, 0).text = "Category"
        t1.cell(0, 1).text = "Traditional"
        t1.cell(0, 2).text = "Hybrid Mode"
        t1.cell(1, 0).text = "Delivery"
        t1.cell(1, 1).text = "Classroom Only"
        t1.cell(1, 2).text = "Blended Classroom + Digital"
        t1.cell(2, 0).text = "Flexibility"
        t1.cell(2, 1).text = "Low"
        t1.cell(2, 2).text = "High"
        doc.add_paragraph("Table 1: Comparison of Educational Delivery Modes")

        # Table 2
        t2 = doc.add_table(rows=3, cols=2)
        t2.cell(0, 0).text = "Tool Type"
        t2.cell(0, 1).text = "Primary Function"
        t2.cell(1, 0).text = "LMS"
        t2.cell(1, 1).text = "Course Content Management"
        t2.cell(2, 0).text = "Video Conferencing"
        t2.cell(2, 1).text = "Synchronous Interaction"
        doc.add_paragraph("Table 2: ICT Tools in Hybrid Mode")

        doc.save(self.docx_path)
        self.amber_template_zip = "C:/Users/shari/Downloads/Basic_book_template__by_Amber_Jain_.zip"

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_amber_jain_generic_book_conversion(self):
        self.assertTrue(os.path.exists(self.amber_template_zip), "Amber Jain template zip missing!")
        
        # 1. Parse Source DOCX
        udm = DocxParser.parse(self.docx_path)
        
        # Verify UDM metadata
        self.assertEqual(udm.metadata.title, "Role of Technology in a Hybrid Teaching Mode")
        self.assertEqual(len(udm.metadata.authors), 1, f"Expected 1 author, got {len(udm.metadata.authors)}: {[a.name for a in udm.metadata.authors]}")
        self.assertEqual(udm.metadata.authors[0].name, "SHARIKA T R")
        
        # Verify no section titles or phone numbers became authors
        author_names = [a.name for a in udm.metadata.authors]
        self.assertNotIn("Phone Number: 9562799289", author_names)
        self.assertNotIn("Introduction", author_names)
        self.assertNotIn("Hybrid vs Blended Learning", author_names)
        self.assertNotIn("Conclusion", author_names)
        
        # 2. Analyze Destination Template
        template_extract_dir = os.path.join(self.test_dir, "template_extracted")
        import zipfile
        with zipfile.ZipFile(self.amber_template_zip, "r") as zf:
            zf.extractall(template_extract_dir)
            
        spec = TemplateAnalyzer.analyze_destination_template(template_extract_dir)
        
        # 3. Render Output LaTeX Project
        output_dir = os.path.join(self.test_dir, "output")
        output_zip = os.path.join(self.test_dir, "converted_project.zip")
        
        created_files = LatexRenderer.render_project(udm, spec, template_extract_dir, output_dir, output_zip)
        
        main_tex_path = os.path.join(output_dir, "main.tex")
        self.assertTrue(os.path.exists(main_tex_path))
        
        with open(main_tex_path, "r", encoding="utf-8") as fh:
            main_tex = fh.read()
            
        print("\n--- GENERATED MAIN.TEX FOR AMBER JAIN TEMPLATE ---")
        print(main_tex)
        print("--------------------------------------------------")
        
        # VERIFICATION ASSERTIONS:
        # 1. Exactly one title
        self.assertIn("\\title{Role of Technology in a Hybrid Teaching Mode}", main_tex)
        self.assertNotIn("Sample Book Title", main_tex)
        self.assertNotIn("Sample book subtitle", main_tex)
        
        # 2. Exactly one author
        self.assertIn("SHARIKA T R", main_tex)
        self.assertNotIn("First-name Last-name", main_tex)
        self.assertNotIn("Phone Number", main_tex)
        self.assertNotIn("author{Introduction", main_tex)
        self.assertNotIn("author{Conclusion", main_tex)
        
        # 3. No template sample content
        self.assertNotIn("Calvin and Hobbes", main_tex)
        self.assertNotIn("Lorem ipsum", main_tex)
        self.assertNotIn("Amber Jain", main_tex)
        
        # 4. Book document hierarchy mapped cleanly to \chapter and \section
        self.assertIn("\\chapter{Introduction}", main_tex)
        self.assertIn("\\chapter{Hybrid vs Blended Learning}", main_tex)
        self.assertIn("\\chapter{ICT and Hybrid Education}", main_tex)
        self.assertIn("\\chapter{Benefits of Hybrid Learning}", main_tex)
        self.assertIn("\\chapter{Challenges in Hybrid Education}", main_tex)
        self.assertIn("\\chapter{Conclusion}", main_tex)
        
        # 5. Source paragraphs preserved
        self.assertIn("Hybrid teaching integrates face-to-face instruction", main_tex)
        self.assertIn("Information and Communication Technologies form the backbone", main_tex)
        
        # 6. Both tables preserved
        self.assertIn("\\begin{table}", main_tex)
        self.assertIn("Traditional", main_tex)
        self.assertIn("LMS", main_tex)
        
        # 7. Test PDF compilation
        output_pdf = os.path.join(self.test_dir, "preview.pdf")
        success, log = LatexSandbox.compile_project(output_dir, "main.tex", udm, output_pdf)
        print("Compilation result status:", "SUCCESS" if success else "FAILED")
        if not success:
            print("LaTeX log:", log)
        self.assertTrue(success, f"PDF compilation failed:\n{log}")


if __name__ == "__main__":
    unittest.main()
