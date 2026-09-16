import os
import sys
import docx

sys.path.insert(0, os.path.abspath(os.path.dirname(__file__)))
from app.parsers.docx_parser import DocxParser
from app.template_engine.analyzer import TemplateAnalyzer
from app.mappers.mapping_engine import MappingEngine
from app.renderers.docx_renderer import DocxRenderer

def run_cdm_regression_tests():
    src_path = os.path.join(os.path.dirname(__file__), "tests", "fixtures", "cdm", "CDM_Theory.docx")
    tpl_path = os.path.join(os.path.dirname(__file__), "tests", "fixtures", "cdm", "WP_NEWFORMAT_CDM_Theory.docx")
    out_dir = os.path.join(os.path.dirname(__file__), "tests", "fixtures", "cdm", "output")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "transformed_cdm_result.docx")
    
    assert os.path.exists(src_path), f"Source CDM missing: {src_path}"
    assert os.path.exists(tpl_path), f"Target CDM template missing: {tpl_path}"
    
    print("\n=== RUNNING CDM TRANSFORMATION REGRESSION SUITE ===")
    
    # 1. Parse Source Document into UDM
    udm = DocxParser.parse(src_path)
    assert udm.doc_type == "Course Document", f"Expected 'Course Document', got {udm.doc_type}"
    assert len(udm.metadata.label_values) > 10, f"Expected >10 label values, got {len(udm.metadata.label_values)}"
    assert len(udm.sections) >= 5, f"Expected >=5 sections, got {len(udm.sections)}"
    assert len(udm.signatures) >= 1, f"Expected signatures detected, got {len(udm.signatures)}"
    print("[PASS] STEP 1: Source CDM parsed into Universal Document Model cleanly")
    
    # 2. Analyze Destination Template
    spec = TemplateAnalyzer.analyze_destination_template(tpl_path)
    assert spec.format_type == "docx"
    assert spec.doc_type_hint == "Course Document"
    assert len(spec.template_tables) >= 5, f"Expected >=5 template tables, got {len(spec.template_tables)}"
    assert len(spec.field_labels) > 5, f"Expected field labels in template, got {len(spec.field_labels)}"
    print("[PASS] STEP 2: Destination CDM Template analyzed into TemplateSpecification")
    
    # 3. Evaluate Mapping Engine
    mapping_res = MappingEngine.map_and_evaluate(udm, spec)
    assert mapping_res["compatibility"]["overall"] >= 65.0, f"Expected compatibility >= 65%, got {mapping_res['compatibility']['overall']}"
    assert mapping_res["label_values_mapped"] > 0, "Expected mapped label-values"
    print(f"[PASS] STEP 3: Content-to-Template Mapping Engine evaluated (Score: {mapping_res['compatibility']['overall']}%, Confidence: {mapping_res['compatibility']['confidence_level']})")
    
    # 4. Render Destination Document In-Place
    DocxRenderer.render(udm, spec, out_path, template_path=tpl_path)
    assert os.path.exists(out_path), "Output DOCX file was not generated"
    assert os.path.getsize(out_path) > 10000, "Output DOCX file is too small"
    
    # 5. Inspect Rendered Document Integrity
    out_doc = docx.Document(out_path)
    assert len(out_doc.tables) >= len(spec.template_tables), "Target template table structure collapsed"
    
    # Verify field values populated inside target table 1
    t1_cells = [c.text for row in out_doc.tables[0].rows for c in row.cells]
    t1_text = " ".join(t1_cells)
    assert len(t1_text) > 100, "Target template table 1 is empty"
    print("[PASS] STEP 4: In-place DOCX rendering & layout preservation verified")
    
    print("\n==================================================")
    print("ALL CDM REGRESSION TESTS PASSED 100%!")
    print("CDM SOURCE -> TARGET TEMPLATE TRANSFORMATION SUCCEEDED!")
    print("==================================================\n")

if __name__ == "__main__":
    run_cdm_regression_tests()
