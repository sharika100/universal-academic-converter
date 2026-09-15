import os
import zipfile
import docx
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw

def generate_samples(base_dir: str):
    os.makedirs(base_dir, exist_ok=True)
    
    # 1. Create IEEE Source Project ZIP
    ieee_dir = os.path.join(base_dir, "ieee_temp")
    os.makedirs(os.path.join(ieee_dir, "figures"), exist_ok=True)
    
    main_tex = r"""\documentclass[conference]{IEEEtran}
\usepackage{cite}
\usepackage{amsmath,amssymb,amsfonts}
\usepackage{graphicx}
\usepackage{textcomp}
\usepackage{xcolor}

\begin{document}

\title{Universal Academic Document Compiler: A Unified Intermediate Representation Approach}

\author{\IEEEauthorblockN{Alice Smith}
\IEEEauthorblockA{\textit{Department of Computer Science} \\
\textit{Stanford University}\\
Stanford, CA, USA \\
alice@stanford.edu}
\and
\IEEEauthorblockN{Bob Jones}
\IEEEauthorblockA{\textit{School of Interactive Computing} \\
\textit{Georgia Institute of Technology}\\
Atlanta, GA, USA \\
bjones@gatech.edu}
}

\maketitle

\begin{abstract}
Converting academic manuscripts between publisher formats (such as IEEE, Springer, Elsevier, and ACM) is a tedious and error-prone manual task for researchers. In this paper, we propose a Universal Document Model (UDM) compiler architecture that decouples source manuscript parsing from destination target rendering.
\end{abstract}

\begin{IEEEkeywords}
document compiler, LaTeX conversion, universal document model, academic publishing
\end{IEEEkeywords}

\section{Introduction}
Academic manuscript conversion traditionally relies on manual reformatting or brittle pair-wise converters. Our compiler-based approach parses input manuscripts into an abstract syntax tree representing semantic document elements.

\section{Methodology}
The system architecture consists of a source parser, universal intermediate representation, destination template analyzer, and target renderer.

\subsection{Universal Document Model}
The model preserves mathematical equations such as:
\begin{equation}
E(m, c) = m \cdot c^2
\label{eq:einstein}
\end{equation}

And tabular experimental results as shown in Table~\ref{tab:results}.

\begin{table}[htbp]
\caption{Performance Comparison across Document Workflows}
\label{tab:results}
\begin{center}
\begin{tabular}{|c|c|c|}
\hline
\textbf{Workflow} & \textbf{Parsing Acc.} & \textbf{Conformity} \\
\hline
DOCX to DOCX & 98.2\% & 97.5\% \\
LaTeX to LaTeX ZIP & 96.8\% & 98.1\% \\
DOCX to LaTeX ZIP & 95.4\% & 96.9\% \\
\hline
\end{tabular}
\end{center}
\end{table}

\begin{figure}[htbp]
\centering
\includegraphics[width=0.8\linewidth]{figures/architecture.png}
\caption{System architecture of the Universal Academic Format Converter.}
\label{fig:arch}
\end{figure}

\section{Conclusion}
We have presented a unified document compiler framework that preserves scientific integrity while ensuring high target template fidelity.

\bibliographystyle{IEEEtran}
\bibliography{references}

\end{document}
"""
    with open(os.path.join(ieee_dir, "main.tex"), "w", encoding="utf-8") as fh:
        fh.write(main_tex)
        
    bib_tex = r"""@article{smith2024universal,
  title={Universal Document Compiler for Academic Publishing},
  author={Smith, Alice and Jones, Bob},
  journal={IEEE Transactions on Knowledge and Data Engineering},
  volume={36},
  number={4},
  pages={1200--1212},
  year={2024},
  publisher={IEEE}
}
@inproceedings{jones2023latex,
  title={Automated LaTeX Template Structural Mapping},
  author={Jones, Bob and Davis, Clara},
  booktitle={Proceedings of ACM International Conference on Document Engineering},
  pages={45--56},
  year={2023}
}
"""
    with open(os.path.join(ieee_dir, "references.bib"), "w", encoding="utf-8") as fh:
        fh.write(bib_tex)
        
    # Create dummy PNG figure
    img = Image.new('RGB', (400, 200), color = (30, 41, 59))
    d = ImageDraw.Draw(img)
    d.text((40, 90), "UDM System Architecture Diagram", fill=(255,255,255))
    img.save(os.path.join(ieee_dir, "figures", "architecture.png"))
    
    # Zip IEEE source project
    ieee_zip_path = os.path.join(base_dir, "ieee_paper.zip")
    with zipfile.ZipFile(ieee_zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for root, _, files in os.walk(ieee_dir):
            for f in files:
                abs_p = os.path.join(root, f)
                rel_p = os.path.relpath(abs_p, ieee_dir)
                zf.write(abs_p, rel_p)
                
    # 2. Create Springer Destination Template ZIP
    springer_dir = os.path.join(base_dir, "springer_temp")
    os.makedirs(springer_dir, exist_ok=True)
    
    with open(os.path.join(springer_dir, "sn-jnl.cls"), "w", encoding="utf-8") as fh:
        fh.write("% Springer Nature Journal Class File (sn-jnl.cls)\n\\NeedsTeXFormat{LaTeX2e}\n\\ProvidesClass{sn-jnl}[2024/01/01 v1.0]\n\\LoadClass{article}\n")
        
    with open(os.path.join(springer_dir, "sn-bibliography.bst"), "w", encoding="utf-8") as fh:
        fh.write("% Springer Nature Bibliography Style File\nENTRY { author title journal year }\n")
        
    springer_sample = r"""\documentclass[pdflatex,sn-mathphys]{sn-jnl}
\usepackage{graphicx}
\usepackage{amsmath,amssymb}

\begin{document}
\title[Short Title]{Springer Template Manuscript Title}
\author[1]{\fnm{FirstName} \sur{LastName}}\email{author@springer.com}
\affiliation[1]{\orgname{Springer Research Center}, \orgaddress{\city{Berlin}, \country{Germany}}}

\abstract{Sample Springer journal manuscript abstract.}
\keywords{Springer, LaTeX template, journal}

\section{Section Header}\label{sec1}
Sample text for Springer format.

\bibliographystyle{sn-bibliography}
\bibliography{sn-bibliography}
\end{document}
"""
    with open(os.path.join(springer_dir, "sample.tex"), "w", encoding="utf-8") as fh:
        fh.write(springer_sample)
        
    springer_zip_path = os.path.join(base_dir, "springer_template.zip")
    with zipfile.ZipFile(springer_zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for root, _, files in os.walk(springer_dir):
            for f in files:
                abs_p = os.path.join(root, f)
                rel_p = os.path.relpath(abs_p, springer_dir)
                zf.write(abs_p, rel_p)

    # 3. Create Multi-Figure and Multi-Author Sample DOCX Manuscript
    docx_doc = docx.Document()
    docx_doc.add_heading("Universal Academic Document Model Architecture", level=0)
    docx_doc.add_paragraph("Alice Smith1, Bob Jones1,2, Charlie Brown2")
    docx_doc.add_paragraph("1 Department of Computer Science, Stanford University, CA, USA")
    docx_doc.add_paragraph("2 Department of Electrical Engineering, MIT, MA, USA")
    docx_doc.add_paragraph("Abstract: Academic document conversion requires semantic structure and multi-figure preservation.")
    docx_doc.add_heading("1. Introduction", level=1)
    docx_doc.add_paragraph("This manuscript demonstrates automatic document compilation from DOCX format.")
    
    # Create 3 distinct images for sample DOCX
    img1_path = os.path.join(base_dir, "tmp_sample_fig1.png")
    img2_path = os.path.join(base_dir, "tmp_sample_fig2.png")
    img3_path = os.path.join(base_dir, "tmp_sample_fig3.png")
    
    i1 = Image.new('RGB', (350, 150), color=(220, 38, 38))
    ImageDraw.Draw(i1).text((40, 60), "Figure 1: Architecture Pipeline", fill=(255,255,255))
    i1.save(img1_path)
    
    i2 = Image.new('RGB', (350, 150), color=(16, 185, 129))
    ImageDraw.Draw(i2).text((40, 60), "Figure 2: Data Flow Chart", fill=(255,255,255))
    i2.save(img2_path)
    
    i3 = Image.new('RGB', (350, 150), color=(37, 99, 235))
    ImageDraw.Draw(i3).text((40, 60), "Figure 3: Experimental Results", fill=(255,255,255))
    i3.save(img3_path)
    
    docx_doc.add_paragraph("Overview of system design is depicted below.")
    docx_doc.add_picture(img1_path, width=Inches(3))
    docx_doc.add_paragraph("Figure 1: Architecture Pipeline Diagram")
    
    docx_doc.add_heading("2. System Design", level=1)
    docx_doc.add_paragraph("The parser extracts sections, tables, and figures.")
    docx_doc.add_picture(img2_path, width=Inches(3))
    docx_doc.add_paragraph("Figure 2: Data Flow Chart")
    
    docx_doc.add_heading("3. Experimental Evaluation", level=1)
    docx_doc.add_paragraph("Performance across workflows is shown in Figure 3.")
    docx_doc.add_picture(img3_path, width=Inches(3))
    docx_doc.add_paragraph("Figure 3: Experimental Results Graph")
    
    t = docx_doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Metric"
    t.cell(0, 1).text = "Score"
    t.cell(1, 0).text = "Parsing Confidence"
    t.cell(1, 1).text = "97%"
    
    docx_doc.add_heading("References", level=1)
    docx_doc.add_paragraph("[1] Smith, A. Universal Document Compiler, 2024.")
    
    docx_sample_path = os.path.join(base_dir, "sample_manuscript.docx")
    docx_doc.save(docx_sample_path)
    
    # Cleanup temporary images
    for p in [img1_path, img2_path, img3_path]:
        if os.path.exists(p):
            os.remove(p)
            
    print("Sample packages created successfully with multi-figure and multi-author data.")

if __name__ == "__main__":
    generate_samples("C:/Users/shari/.gemini/antigravity/scratch/universal-academic-converter/samples")
