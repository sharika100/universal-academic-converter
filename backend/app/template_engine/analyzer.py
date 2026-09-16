import os
import re
import docx
from typing import Tuple, List, Dict, Any
from app.models.template_spec import TemplateSpecification
from app.parsers.zip_utils import find_latex_entrypoint

class TemplateAnalyzer:
    @staticmethod
    def analyze_destination_template(template_path: str) -> TemplateSpecification:
        """Analyzes a target DOCX template file or LaTeX ZIP template project."""
        if template_path.endswith(".docx"):
            return TemplateAnalyzer._analyze_docx_template(template_path)
        elif template_path.endswith(".zip") or os.path.isdir(template_path):
            return TemplateAnalyzer._analyze_latex_zip_template(template_path)
        else:
            spec = TemplateSpecification(format_type="latex")
            spec.detected_rules.append("Single TeX/CLS template file uploaded")
            return spec

    @staticmethod
    def _analyze_docx_template(docx_path: str) -> TemplateSpecification:
        spec = TemplateSpecification(format_type="docx", document_class="Word Template")
        doc = docx.Document(docx_path)
        
        detected_styles = [s.name for s in doc.styles]
        spec.detected_rules.append(f"Detected {len(detected_styles)} DOCX paragraph/character styles")
        
        # Document classification hint
        full_text = " ".join([p.text for p in doc.paragraphs[:30]]).lower()
        if any(kw in full_text for kw in ["course delivery manual", "cdm", "syllabus", "course outcome"]):
            spec.doc_type_hint = "Course Document"
        elif any(kw in full_text for kw in ["requisition form", "application form", "approval form"]):
            spec.doc_type_hint = "Institutional Form"
        else:
            spec.doc_type_hint = "General Document"

        # 1. Page Setup Analysis
        if doc.sections:
            sec = doc.sections[0]
            spec.page_setup = {
                "page_width": sec.page_width.inches if sec.page_width else 8.5,
                "page_height": sec.page_height.inches if sec.page_height else 11.0,
                "top_margin": sec.top_margin.inches if sec.top_margin else 1.0,
                "bottom_margin": sec.bottom_margin.inches if sec.bottom_margin else 1.0,
                "left_margin": sec.left_margin.inches if sec.left_margin else 1.0,
                "right_margin": sec.right_margin.inches if sec.right_margin else 1.0,
                "orientation": str(sec.orientation) if hasattr(sec, "orientation") else "PORTRAIT"
            }

        # 2. Table Specifications Extraction
        spec_tables = []
        field_labels = []
        signature_slots = []
        
        for t_idx, tbl in enumerate(doc.tables):
            rows = len(tbl.rows)
            cols = len(tbl.columns) if tbl.rows else 0
            fixed_labels = []
            editable_slots = []
            is_signature_table = False
            
            for r_idx, row in enumerate(tbl.rows):
                cell_txts = [c.text.strip() for c in row.cells]
                for c_idx, txt in enumerate(cell_txts):
                    if any(kw in txt.lower() for kw in ["signature", "hod", "course instructor", "stream coordinator"]):
                        is_signature_table = True
                        
                    if txt.endswith(":") or any(kw in txt.lower() for kw in ["name", "offering", "semester", "instructor", "sl no", "co no", "code", "description", "questions"]):
                        fixed_labels.append({"row": r_idx, "col": c_idx, "text": txt})
                        field_labels.append(txt)
                    elif not txt or txt.startswith(":") or re.match(r'^[\{\[\<]', txt):
                        editable_slots.append({"row": r_idx, "col": c_idx, "initial": txt})
                        
            if is_signature_table:
                signature_slots.append({"table_index": t_idx, "rows": rows, "cols": cols})
                
            spec_tables.append({
                "table_index": t_idx,
                "rows": rows,
                "cols": cols,
                "fixed_labels": fixed_labels,
                "editable_slots": editable_slots
            })
            
        spec.template_tables = spec_tables
        spec.field_labels = field_labels
        spec.signature_slots = signature_slots
        spec.template_confidence = 98.0
        return spec

    @staticmethod
    def _analyze_latex_zip_template(project_dir: str) -> TemplateSpecification:
        spec = TemplateSpecification(format_type="latex")
        rules = []
        warnings = []
        required_files = []
        
        for root, _, files in os.walk(project_dir):
            for f in files:
                rel_f = os.path.relpath(os.path.join(root, f), project_dir).replace("\\", "/")
                ext = os.path.splitext(f)[1].lower()
                if ext in [".cls", ".sty", ".bst"] or f.endswith(".png") or f.endswith(".eps") or f.endswith(".pdf"):
                    required_files.append(rel_f)
                    
        spec.required_files = required_files
        rules.append(f"Detected {len(required_files)} target template infrastructure files (.cls, .sty, .bst, assets)")
        
        entrypoint_rel, candidates, _ = find_latex_entrypoint(project_dir)
        if entrypoint_rel:
            spec.entry_point_file = entrypoint_rel
            main_p = os.path.join(project_dir, entrypoint_rel)
            try:
                with open(main_p, "r", encoding="utf-8", errors="ignore") as fh:
                    sample_content = fh.read()
                    spec.sample_content = sample_content
                    
                cls_match = re.search(r'\\documentclass(?:\[([^\]]*)\])?\{([^}]+)\}', sample_content)
                if cls_match:
                    opts_str = cls_match.group(1)
                    spec.document_class = cls_match.group(2).strip()
                    if opts_str:
                        spec.class_options = [o.strip() for o in opts_str.split(",")]
                    rules.append(f"Document class: \\documentclass[{opts_str or ''}]{{{spec.document_class}}}")
                    
                # Determine Author Formatting Style
                if "IEEEauthorblockN" in sample_content or "IEEEtran" in spec.document_class:
                    spec.author_style = "ieee"
                    rules.append("Author block: IEEE (\\IEEEauthorblockN & \\IEEEauthorblockA)")
                elif r"\fnm" in sample_content or r"\sur" in sample_content or "sn-jnl" in spec.document_class or r"\affil" in sample_content:
                    spec.author_style = "springer"
                    rules.append("Author block: Springer (\\author[id]{\\fnm{} \\sur{}}, \\affil[id]{})")
                elif "elsarticle" in spec.document_class or r"\address" in sample_content:
                    spec.author_style = "elsevier"
                    rules.append("Author block: Elsevier (\\author[id]{}, \\address[id]{})")
                elif "llncs" in spec.document_class or r"\inst" in sample_content:
                    spec.author_style = "lncs"
                    rules.append("Author block: LNCS (\\author{...\\inst{1}}, \\institute{})")
                elif "acmart" in spec.document_class:
                    spec.author_style = "acm"
                    rules.append("Author block: ACM (\\author{}, \\affiliation{})")
                else:
                    spec.author_style = "standard"
                    rules.append("Author block: Standard LaTeX (\\author{}, \\institute{})")
                    
                bst_match = re.search(r'\\bibliographystyle\{([^}]+)\}', sample_content)
                if bst_match:
                    spec.bib_style = bst_match.group(1).strip()
                    rules.append(f"Bibliography style: \\bibliographystyle{{{spec.bib_style}}}")
                    
                if "natbib" in sample_content:
                    spec.citation_system = "natbib"
                    rules.append("Citation system: natbib")
                elif "biblatex" in sample_content:
                    spec.citation_system = "biblatex"
                    rules.append("Citation system: biblatex")
                else:
                    spec.citation_system = "numeric"
                    rules.append("Citation system: Standard numeric")
                    
            except Exception as e:
                warnings.append(f"Sample file analysis warning: {str(e)}")
        else:
            warnings.append("No sample .tex entrypoint detected in target template zip.")
            
        spec.detected_rules = rules
        spec.warnings = warnings
        spec.template_confidence = 98.0
        return spec
