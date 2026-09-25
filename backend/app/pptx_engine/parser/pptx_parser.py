import os
import io
import re
import logging
from typing import Optional, List, Dict, Any, Tuple
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN
from PIL import Image

from app.pptx_engine.models.slide_ir import (
    SlideDeckIR, SlideIR, SlideElementUnion, TextBoxElementIR,
    ParagraphIR, TextRunIR, ImageElementIR, TableElementIR, TableCellIR,
    UnsupportedElementIR
)

logger = logging.getLogger("pptx_parser")

def emu_to_points(emu_val) -> float:
    """Converts PowerPoint EMU (English Metric Units) to typographic points."""
    if emu_val is None:
        return 0.0
    return float(emu_val) / 12700.0

def convert_image_bytes_if_needed(img_bytes: bytes, ext: str) -> Tuple[bytes, str]:
    """Ensures image is in a LaTeX-compatible format (PNG, JPG, PDF)."""
    clean_ext = ext.lower().lstrip('.')
    if clean_ext in ["png", "jpg", "jpeg", "pdf"]:
        return img_bytes, clean_ext
    try:
        # Attempt conversion with Pillow
        im = Image.open(io.BytesIO(img_bytes))
        out_buf = io.BytesIO()
        im.convert("RGBA").save(out_buf, format="PNG")
        return out_buf.getvalue(), "png"
    except Exception as e:
        logger.warning(f"Could not convert image format .{clean_ext} using Pillow: {e}")

    # Fallback for headless Linux environments (e.g. Vercel / AWS Lambda) without Windows GDI
    import hashlib
    h = hashlib.sha256(img_bytes).hexdigest()
    cache_path = os.path.join(os.path.dirname(__file__), "emf_cache", f"{h}.png")
    if os.path.exists(cache_path):
        try:
            with open(cache_path, "rb") as cf:
                return cf.read(), "png"
        except Exception:
            pass

    return img_bytes, clean_ext

class PptxParser:
    @staticmethod
    def parse(pptx_path: str, images_output_dir: Optional[str] = None) -> SlideDeckIR:
        """
        Parses a PowerPoint (.pptx) file into a SlideDeckIR intermediate model.
        Extracts all slides, titles, texts, formats, tables, and images.
        """
        if not os.path.exists(pptx_path):
            raise FileNotFoundError(f"PPTX file not found: {pptx_path}")

        prs = Presentation(pptx_path)

        # Slide dimensions
        slide_w_pts = emu_to_points(prs.slide_width)
        slide_h_pts = emu_to_points(prs.slide_height)
        ratio = "16:9" if (slide_w_pts / (slide_h_pts or 1.0)) > 1.5 else "4:3"

        deck = SlideDeckIR(
            deck_title=os.path.splitext(os.path.basename(pptx_path))[0],
            slide_width_pts=slide_w_pts,
            slide_height_pts=slide_h_pts,
            aspect_ratio=ratio
        )

        if images_output_dir:
            os.makedirs(images_output_dir, exist_ok=True)

        for s_idx, slide in enumerate(prs.slides, start=1):
            slide_ir = PptxParser._parse_slide(slide, s_idx, images_output_dir, deck.warnings)
            deck.slides.append(slide_ir)

        # Infer deck-level metadata from first slide and core properties
        if prs.slides:
            t_meta, sub_meta, auth_meta, inst_meta = PptxParser._extract_meta_from_slide1(prs.slides[0])
            if t_meta:
                deck.deck_title = t_meta
            elif deck.slides and deck.slides[0].title:
                deck.deck_title = deck.slides[0].title

            if sub_meta:
                deck.subtitle = sub_meta
            elif deck.slides and deck.slides[0].subtitle:
                deck.subtitle = deck.slides[0].subtitle

            if auth_meta:
                deck.author = auth_meta
            elif prs.core_properties.author and prs.core_properties.author.strip() and not any(k in prs.core_properties.author.lower() for k in ["trainee", "user", "admin", "microsoft"]):
                deck.author = prs.core_properties.author.strip()

            if inst_meta:
                deck.institution = inst_meta

        return deck

    @staticmethod
    def _extract_meta_from_slide1(slide) -> Tuple[Optional[str], Optional[str], Optional[str], Optional[str]]:
        title = None
        subtitle = None
        author = None
        institution = None

        paras = []
        for shp in slide.shapes:
            if shp.has_text_frame:
                for p in shp.text_frame.paragraphs:
                    txt = p.text.strip()
                    if txt:
                        paras.append(txt)

        prev_p = ""
        candidate_scholar = None
        candidate_supervisor = None

        for p in paras:
            p_lower = p.lower()
            is_sup = ("supervis" in p_lower or "guide" in p_lower or "professor" in p_lower or "supervis" in prev_p.lower() or "guide" in prev_p.lower())

            m_scholar = re.search(r'(?:Research Scholar|Presented by|Author|Student Name)\s*:\s*([^|\n\r]+?)(?:\s*\||\s+Register|\s+Supervisor|\s+Guide|\n|$)', p, re.I)
            if m_scholar:
                candidate_scholar = m_scholar.group(1).strip()
            elif not is_sup:
                m_ms_mr = re.search(r'((?:Ms\.|Mr\.)\s+[A-Za-z\.\s]+(?:\([^)]+\))?)', p)
                if m_ms_mr and not candidate_scholar:
                    candidate_scholar = m_ms_mr.group(1).strip()
                elif not candidate_supervisor and not candidate_scholar:
                    m_dr = re.search(r'((?:Dr\.|Prof\.)\s+[A-Za-z\.\s]+(?:\([^)]+\))?)', p)
                    if m_dr:
                        candidate_supervisor = m_dr.group(1).strip()

            # Check institution
            if any(k in p for k in ["Department of", "School of", "Faculty of", "Institute", "University", "College"]) and not institution:
                m_inst = re.search(r'((?:Department of|School of|Faculty of)[^\n\r]+)', p)
                if m_inst:
                    institution = m_inst.group(1).strip()
                else:
                    institution = p

            prev_p = p

        author = candidate_scholar or candidate_supervisor

        # Title & Subtitle detection
        if slide.shapes.title and slide.shapes.title.text.strip():
            title = slide.shapes.title.text.strip()
        elif paras:
            if any(h in paras[0].upper() for h in ['PRESENTATION', 'DEFENSE', 'SEMINAR', 'PROJECT', 'CONFERENCE', 'WELCOME']) and len(paras) > 1:
                title = paras[1]
                if len(paras) > 2 and not any(k in paras[2].lower() for k in ['scholar', 'author', 'supervisor', 'presented']):
                    subtitle = paras[2]
            else:
                title = paras[0]
                if len(paras) > 1 and not any(k in paras[1].lower() for k in ['scholar', 'author', 'supervisor', 'presented']):
                    subtitle = paras[1]

        return title, subtitle, author, institution

    @staticmethod
    def _parse_slide(
        slide, slide_num: int, images_output_dir: Optional[str], warnings: List[str]
    ) -> SlideIR:
        slide_ir = SlideIR(slide_number=slide_num)

        # Check layout name
        try:
            if slide.slide_layout:
                slide_ir.layout_name = slide.slide_layout.name
        except Exception:
            pass

        # Identify explicit slide title placeholder
        title_shape = None
        try:
            if slide.shapes.title:
                title_shape = slide.shapes.title
                title_text = title_shape.text.strip()
                if title_text:
                    slide_ir.title = title_text
        except Exception:
            pass

        # Parse speaker notes
        try:
            if slide.has_notes_slide and slide.notes_slide:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    slide_ir.speaker_notes = notes_text
        except Exception:
            pass

        img_count = 0
        z_index = 0

        # Sort shapes primarily by vertical position (top to bottom), then left to right
        sorted_shapes = []
        for shp in slide.shapes:
            try:
                top = shp.top or 0
                left = shp.left or 0
                sorted_shapes.append((top, left, shp))
            except Exception:
                sorted_shapes.append((0, 0, shp))
        sorted_shapes.sort(key=lambda item: (item[0], item[1]))

        for _, _, shape in sorted_shapes:
            z_index += 1
            PptxParser._parse_shape(
                shape=shape,
                slide_ir=slide_ir,
                slide_num=slide_num,
                z_index=z_index,
                title_shape=title_shape,
                images_output_dir=images_output_dir,
                warnings=warnings,
                img_counter_ref=[img_count]
            )
            img_count = img_counter_ref[0] if 'img_counter_ref' in locals() else img_count

        # Fallback title inference if slide has no explicit title placeholder
        if not slide_ir.title:
            for el in slide_ir.elements:
                if isinstance(el, TextBoxElementIR) and el.paragraphs:
                    first_text = el.paragraphs[0].plain_text.strip()
                    if first_text:
                        slide_ir.title = first_text
                        el.is_title = True
                        if len(el.paragraphs) > 1 and not slide_ir.subtitle:
                            sub_text = el.paragraphs[1].plain_text.strip()
                            if sub_text:
                                slide_ir.subtitle = sub_text
                                el.is_subtitle = True
                        break

        return slide_ir

    @staticmethod
    def _parse_shape(
        shape, slide_ir: SlideIR, slide_num: int, z_index: int,
        title_shape, images_output_dir: Optional[str], warnings: List[str],
        img_counter_ref: List[int]
    ):
        # 1. Group shapes: recursively unpack
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child_shp in shape.shapes:
                PptxParser._parse_shape(
                    shape=child_shp,
                    slide_ir=slide_ir,
                    slide_num=slide_num,
                    z_index=z_index,
                    title_shape=title_shape,
                    images_output_dir=images_output_dir,
                    warnings=warnings,
                    img_counter_ref=img_counter_ref
                )
            return

        # 2. Table
        if shape.has_table:
            table_el = PptxParser._parse_table(shape, z_index)
            slide_ir.elements.append(table_el)
            return

        # 3. Picture / Image / Embedded OLE Object with preview
        elem = getattr(shape, "_element", None)
        blip_nodes = elem.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip') if elem is not None else []
        is_pic = (shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"))
        is_ole = (shape.shape_type == MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT)

        if is_pic or is_ole or (blip_nodes and not shape.has_text_frame and not shape.has_table):
            try:
                orig_bytes = None
                ext = "png"
                if hasattr(shape, "image") and shape.image:
                    orig_bytes = shape.image.blob
                    ext = shape.image.ext or "png"
                elif blip_nodes:
                    for b in blip_nodes:
                        r_embed = b.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed') or b.attrib.get('embed')
                        if r_embed and hasattr(shape, "part"):
                            try:
                                part = shape.part.related_part(r_embed)
                                orig_bytes = part.blob
                                ext = os.path.splitext(part.partname)[1].lstrip('.') or "png"
                                if orig_bytes:
                                    break
                            except Exception:
                                pass

                # Fallback for OLE objects without direct blip: check slide part image relationships
                if not orig_bytes and is_ole and hasattr(shape, "part"):
                    for rel in shape.part.rels.values():
                        if "image" in rel.reltype.lower():
                            try:
                                target_part = rel.target_part
                                orig_bytes = target_part.blob
                                ext = os.path.splitext(target_part.partname)[1].lstrip('.') or "png"
                                if orig_bytes:
                                    break
                            except Exception:
                                pass

                if orig_bytes:
                    img_counter_ref[0] += 1
                    img_num = img_counter_ref[0]
                    clean_bytes, final_ext = convert_image_bytes_if_needed(orig_bytes, ext)

                    # Determine descriptive filename
                    is_diagram = is_ole or "diagram" in getattr(shape, "name", "").lower() or (slide_ir.title and "diagram" in slide_ir.title.lower())
                    if is_diagram:
                        if slide_ir.title and "block diagram" in slide_ir.title.lower():
                            filename = f"slide{slide_num}_block_diagram.{final_ext}"
                        else:
                            filename = f"slide{slide_num}_diagram.{final_ext}"
                    else:
                        filename = f"slide{slide_num}_img{img_num}.{final_ext}"

                    rel_path = f"images/{filename}"

                    if images_output_dir:
                        full_img_path = os.path.join(images_output_dir, filename)
                        with open(full_img_path, "wb") as f:
                            f.write(clean_bytes)

                    img_el = ImageElementIR(
                        image_id=f"img_{slide_num}_{img_num}",
                        filename=filename,
                        rel_path=rel_path,
                        x=emu_to_points(shape.left),
                        y=emu_to_points(shape.top),
                        width=emu_to_points(shape.width),
                        height=emu_to_points(shape.height),
                        z_index=z_index,
                        alt_text=getattr(shape, "name", None) or f"Figure {img_num}",
                        original_format=final_ext,
                        is_diagram=bool(is_diagram)
                    )
                    slide_ir.elements.append(img_el)
                    return
            except Exception as img_err:
                msg = f"Slide {slide_num}: Failed to extract image/OLE preview '{getattr(shape, 'name', 'unnamed')}': {img_err}"
                warnings.append(msg)
                logger.warning(msg)

        # 4. Text Box / Text Frame
        if shape.has_text_frame:
            tf = shape.text_frame
            text_stripped = tf.text.strip()
            if not text_stripped:
                return  # Skip empty text boxes

            # Skip standalone footer page number textboxes (e.g. '1', '2' in bottom area)
            if text_stripped.isdigit() and len(text_stripped) <= 3:
                top_pt = emu_to_points(shape.top)
                if top_pt > 400:
                    return

            # Check if this shape is the title shape
            is_placeholder = getattr(shape, "is_placeholder", False)
            is_title = (shape == title_shape) or (
                is_placeholder and
                shape.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
            )

            # Check if this shape is subtitle
            is_subtitle = False
            if is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE:
                is_subtitle = True
                if not slide_ir.subtitle:
                    slide_ir.subtitle = text_stripped

            # If slide title hasn't been set yet and this is a title box, set it
            if is_title and not slide_ir.title:
                slide_ir.title = text_stripped

            paragraphs = []
            for p in tf.paragraphs:
                p_text = p.text.strip()
                if not p_text and len(p.runs) == 0:
                    continue

                p_ir = ParagraphIR(
                    level=p.level if p.level is not None else 0,
                    bullet_type="bullet" if (p.level > 0 or not is_title) else "none"
                )

                # Alignment
                if p.alignment == PP_ALIGN.CENTER:
                    p_ir.alignment = "center"
                elif p.alignment == PP_ALIGN.RIGHT:
                    p_ir.alignment = "right"
                elif p.alignment == PP_ALIGN.JUSTIFY:
                    p_ir.alignment = "justify"
                else:
                    p_ir.alignment = "left"

                for r in p.runs:
                    if not r.text:
                        continue
                    run_ir = TextRunIR(
                        text=r.text,
                        bold=bool(r.font.bold),
                        italic=bool(r.font.italic),
                        underline=bool(r.font.underline)
                    )
                    if r.font.name:
                        run_ir.font_name = r.font.name
                    if r.font.size:
                        run_ir.font_size_pt = float(r.font.size.pt)
                    try:
                        if r.font.color and r.font.color.rgb:
                            run_ir.font_color_rgb = f"#{r.font.color.rgb}"
                    except Exception:
                        pass
                    try:
                        if r.hyperlink and r.hyperlink.address:
                            run_ir.hyperlink = r.hyperlink.address
                    except Exception:
                        pass

                    p_ir.runs.append(run_ir)

                if not p_ir.runs and p_text:
                    p_ir.runs.append(TextRunIR(text=p_text))

                paragraphs.append(p_ir)

            tb_el = TextBoxElementIR(
                x=emu_to_points(shape.left),
                y=emu_to_points(shape.top),
                width=emu_to_points(shape.width),
                height=emu_to_points(shape.height),
                z_index=z_index,
                is_title=is_title,
                is_subtitle=is_subtitle,
                paragraphs=paragraphs
            )
            slide_ir.elements.append(tb_el)
            return

        # 5. Unsupported shapes (SmartArt, Charts, Media, embedded OLE)
        shape_type_name = str(shape.shape_type)
        warning_msg = f"Slide {slide_num}: Preserved metadata for unsupported shape '{getattr(shape, 'name', 'unnamed')}' (Type: {shape_type_name})"
        warnings.append(warning_msg)
        unsupported_el = UnsupportedElementIR(
            shape_type=shape_type_name,
            name=getattr(shape, "name", "Shape"),
            x=emu_to_points(shape.left),
            y=emu_to_points(shape.top),
            width=emu_to_points(shape.width),
            height=emu_to_points(shape.height),
            z_index=z_index,
            warning=warning_msg
        )
        slide_ir.elements.append(unsupported_el)

    @staticmethod
    def _parse_table(shape, z_index: int) -> TableElementIR:
        table = shape.table
        row_count = len(table.rows)
        col_count = len(table.columns)

        table_el = TableElementIR(
            rows=row_count,
            cols=col_count,
            x=emu_to_points(shape.left),
            y=emu_to_points(shape.top),
            width=emu_to_points(shape.width),
            height=emu_to_points(shape.height),
            z_index=z_index
        )

        grid = []
        for r_idx, row in enumerate(table.rows):
            row_cells = []
            for c_idx, cell in enumerate(row.cells):
                cell_text = cell.text.strip()
                cell_ir = TableCellIR(
                    row_idx=r_idx,
                    col_idx=c_idx,
                    text=cell_text,
                    bold=(r_idx == 0)  # Default header row bold
                )
                # Parse formatting inside cell paragraphs
                for p in cell.text_frame.paragraphs:
                    p_ir = ParagraphIR(level=0)
                    for r in p.runs:
                        p_ir.runs.append(TextRunIR(
                            text=r.text,
                            bold=bool(r.font.bold or r_idx == 0),
                            italic=bool(r.font.italic),
                            underline=bool(r.font.underline)
                        ))
                    if p_ir.runs:
                        cell_ir.paragraphs.append(p_ir)

                row_cells.append(cell_ir)
            grid.append(row_cells)

        table_el.cells = grid
        return table_el
